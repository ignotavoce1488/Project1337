"""Shared DOCX rendering without temporary files or framework dependencies."""

from io import BytesIO

from docx import Document

from slovech.core.models import Lecture


def render_docx(lecture: Lecture, lang: str = "ru") -> tuple[bytes, str]:
    translated = lang == "ru" and lecture.language == "en"
    title = (lecture.title_ru or lecture.title) if translated else lecture.title
    summary = (lecture.summary_ru or lecture.summary) if translated else lecture.summary
    points = (lecture.key_points_ru or lecture.key_points) if translated else lecture.key_points
    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(f"Дата: {lecture.created_at}")
    document.add_heading("Главные мысли", 1)
    for point in points:
        document.add_paragraph(point.replace("**", ""), style="List Bullet")
    document.add_heading("Конспект", 1)
    for line in summary.splitlines():
        line = line.strip()
        if line.startswith("#"):
            document.add_heading(line.lstrip("# "), 2)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].replace("**", ""), style="List Bullet")
        elif line:
            document.add_paragraph(line.replace("**", ""))
    document.add_heading("Полная расшифровка", 1)
    for line in lecture.transcription.splitlines():
        if line.strip():
            document.add_paragraph(line.strip())
    output = BytesIO()
    document.save(output)
    name = "".join(c for c in title if c.isalnum() or c in " _-").strip()[:60] or "Конспект"
    return output.getvalue(), f"{name}.docx"
